"""Enterprise document loaders that normalize files into LangChain Documents."""

from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document

from utils.file_utils import normalize_file_type, read_text_with_fallback, safe_filename


SUPPORTED_EXTENSIONS = {"pdf", "md", "markdown", "docx", "xlsx", "xls", "txt"}


def _base_metadata(path: Path, file_type: str) -> dict:
    return {
        "source": path.name,
        "file_path": str(path),
        "file_type": file_type,
    }


def _load_text(path: Path, file_type: str) -> list[Document]:
    return [Document(page_content=read_text_with_fallback(path), metadata=_base_metadata(path, file_type))]


def _load_pdf(path: Path) -> list[Document]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("PDF 加载需要安装 pypdf，请执行：pip install pypdf") from exc

    reader = PdfReader(str(path))
    docs = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        metadata = _base_metadata(path, "pdf")
        metadata["page"] = index
        docs.append(Document(page_content=text, metadata=metadata))
    return docs


def _load_docx(path: Path) -> list[Document]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ImportError("Word docx 加载需要安装 python-docx，请执行：pip install python-docx") from exc

    docx = DocxDocument(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in docx.paragraphs if paragraph.text.strip()]
    tables = []
    for table in docx.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                tables.append(" | ".join(cells))
    text = "\n".join([*paragraphs, *tables])
    return [Document(page_content=text, metadata=_base_metadata(path, "docx"))]


def _load_excel(path: Path, file_type: str) -> list[Document]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise ImportError("Excel 加载需要安装 pandas/openpyxl/xlrd，请执行：pip install pandas openpyxl xlrd") from exc

    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    docs = []
    for sheet_name, frame in sheets.items():
        frame = frame.fillna("")
        rows = []
        for _, row in frame.iterrows():
            values = [str(value).strip() for value in row.tolist() if str(value).strip()]
            if values:
                rows.append(" | ".join(values))
        metadata = _base_metadata(path, file_type)
        metadata["sheet_name"] = str(sheet_name)
        docs.append(Document(page_content="\n".join(rows), metadata=metadata))
    return docs


def load_document(path: str | Path) -> list[Document]:
    """Load one local file and return LangChain Document objects."""
    file_path = Path(path)
    file_type = normalize_file_type(file_path.name)
    if file_type == "markdown":
        file_type = "md"
    if file_type not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"暂不支持的文件类型：{file_type}")

    if file_type == "pdf":
        return _load_pdf(file_path)
    if file_type == "docx":
        return _load_docx(file_path)
    if file_type in {"xlsx", "xls"}:
        return _load_excel(file_path, file_type)
    return _load_text(file_path, file_type)


def save_uploaded_file(uploaded_file, upload_dir: str | Path) -> Path:
    """Persist a Streamlit UploadedFile so loaders can use file-path based parsers."""
    target_dir = Path(upload_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    filename = safe_filename(uploaded_file.name)
    target_path = target_dir / filename
    target_path.write_bytes(uploaded_file.getvalue())
    return target_path


def load_uploaded_file(uploaded_file, upload_dir: str | Path) -> list[Document]:
    return load_document(save_uploaded_file(uploaded_file, upload_dir))


def merge_document_text(documents: Iterable[Document]) -> str:
    return "\n\n".join(document.page_content for document in documents)
